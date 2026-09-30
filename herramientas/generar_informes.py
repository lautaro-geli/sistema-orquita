from pathlib import Path
import csv,json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
R=Path(__file__).resolve().parents[1]
M=json.loads((R/'evidencias/metricas-modelo.json').read_text(encoding='utf-8'))
ROWS=list(csv.DictReader((R/'evidencias/predicciones.csv').open(encoding='utf-8'),delimiter=';'))

def base(title,subtitle):
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
 for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  st=d.styles[n];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
  st.paragraph_format.space_after=Pt(7)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.08
 d.styles['Title'].font.size=Pt(23);d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 2'].font.size=Pt(13)
 h=sec.header.paragraphs[0];h.text='Proyecto Orquitas   |   Hito 2 simulado';h.style='Caption'
 for run in h.runs:run.font.color.rgb=RGBColor(0,0,0)
 f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 f.add_run('Página ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 d.add_paragraph(title,'Title');d.add_paragraph(subtitle,'Subtitle')
 return d

def p(d,text):return d.add_paragraph(text)
def table(d,heads,rows,widths):
 t=d.add_table(rows=1, cols=len(heads));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,h in enumerate(heads):t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=str(x)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');b.append(el)
   pr.append(b);mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F5F7FA' if ri%2==0 else 'FFFFFF'));pr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.0
    if ci>0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.size=Pt(10);run.bold=(ri==0)
 d.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

d=base('Ejemplos y resultados de la simulación','Validación de Neuroph y comunicación TCP   •   30 de septiembre de 2026')
p(d,'Este informe documenta una demostración del servidor Orquitas con lecturas ficticias. La red obtuvo 97,80 % de aciertos en 1000 muestras de evaluación posterior y 90,625 % en 4000 lecturas cercanas al umbral. Las seis escenas del simulador produjeron las decisiones previstas. Estos porcentajes no representan pruebas de un sensor real ni éxito en campo.')
d.add_heading('Qué se simuló',1)
p(d,'Se usa un sensor virtual de 12 distancias en milímetros, rango 0 a 4000, sobre una cobertura convencional de 180 grados. Los centros son 7,5° + 15° × índice. Los sectores 4 a 7 forman el frente; 0 a 5, la derecha; 6 a 11, la izquierda. El valor cero significa proximidad en esta simulación. No se seleccionó ni calibró hardware; un cero de un sensor físico podría significar ausencia de eco y requerir otro tratamiento.')
p(d,'LIBRE significa que ningún sector frontal es menor a 400 mm; OBJETO significa que al menos uno lo es. Esta regla solo genera las etiquetas de referencia. En ejecución, el servidor obtiene la clase mediante Neuroph. La red no distingue todavía paredes, lobitos ni salida.')
d.add_heading('Entrenamiento y evaluación',1)
p(d,'Se mantuvo una MLP 12–6–2 con activación sigmoid, entradas divididas por 4000 y MomentumBackpropagation. Se usaron 2400 ejemplos balanceados, semilla 101, pesos con semilla 20260930, tasa 0,1 y momentum 0,7. El generador incluye paredes laterales, objetos frontales estrechos, proximidad general, espacio libre y casos de frontera con ruido de ±5 %.')
p(d,'El entrenamiento terminó en el límite de 4000 iteraciones con error 0,00552914. No alcanzó el objetivo de error 0,003. La semilla 202 y la frontera con semilla 7 se usaron durante desarrollo; no se presentan como evaluación final ciega. Después de fijar el modelo se evaluaron las semillas 424242 y 987654. Los vectores y predicciones se conservan en evidencias/predicciones.csv.')
d.add_heading('Cómo se calculó el porcentaje',1)
p(d,'Exactitud = 100 × cantidad de predicciones correctas / cantidad total de lecturas. Se compara la clase predicha con la etiqueta sintética; no se mide distancia recorrida, capturas físicas ni calidad del sensor.')
d.add_page_break()
d.add_heading('Resultados medidos',1)
table(d,['Conjunto','Muestras','Aciertos','Exactitud'],[
 ['Desarrollo general 202',1000,980,'98,00 %'],
 ['Evaluación posterior 424242',1000,978,'97,80 %'],
 ['Frontera de desarrollo 7',4000,3585,'89,625 %'],
 ['Frontera posterior 987654',4000,3625,'90,625 %'],
 ['Seis escenas de demostración',6,6,'100,00 %']], [3.2,1.05,1.05,1.7])
p(d,'Las diez semillas adicionales 1000–1009 dieron una media de 97,25 %, con mínimo de 96,50 % y máximo de 98,20 %. Pertenecen a la misma familia sintética; cambiar la semilla no equivale a probar otra realidad física.')
d.add_heading('Errores de la evaluación posterior',2)
table(d,['Conjunto','OBJETO clasificado LIBRE','LIBRE clasificado OBJETO'],[
 ['General 424242',2,20],['Frontera 987654',86,289]],[2.7,2.15,2.15])
p(d,'En la frontera, 86 lecturas con objeto se clasificaron como libres. Por eso no corresponde afirmar que todos los errores sean conservadores ni que la red garantice evitar choques. Los 289 falsos objetos también pueden causar giros o detenciones innecesarias.')
d.add_heading('Límite de la evidencia',2)
p(d,'Las etiquetas se construyen con una regla conocida. Aplicar esa regla directamente daría 100 % sobre sus propias etiquetas; esta red demuestra entrenamiento, inferencia e integración TCP, no superioridad sobre la regla. La simulación no demuestra identificación de objetos reales. La precisión cerca de 400 mm debe informarse por separado y no ocultarse detrás del porcentaje general.')
d.add_page_break()
d.add_heading('Seis lecturas de ejemplo',1)
p(d,'Cada vector lista los sectores d0 a d11, en milímetros. El servidor recibe RECOPILACION, infiere la clase y devuelve ORDEN por el mismo socket. Las clases y órdenes de la tabla se obtuvieron con el modelo guardado.')
examples=[x for x in ROWS if x['grupo']=='ejemplo']
names=['Espacio libre','Pared lateral y frente libre','Objeto con espacio a izquierda','Objeto con espacio a derecha','Entorno muy cercano','Espacio libre después de capturas']
table(d,['Caso','Esperado','Red','Orden'],[[str(i+1)+' '+names[i],x['esperado'],x['predicho'],x['orden']] for i,x in enumerate(examples)],[3.05,1,1,1.95])
for i,x in enumerate(examples):
 para=d.add_paragraph();para.add_run(str(i+1)+'  ').bold=True
 vector='['+', '.join(str(int(float(x['d'+str(j)]))) for j in range(12))+']'
 run=para.add_run(vector);run.font.name='Consolas';run.font.size=Pt(10)
p(d,'LIBRE → AVANZAR. OBJETO → girar hacia el lado con mayor suma de distancias; empate a izquierda. Si ambas sumas son menores a 1200 mm → DETENER. Las escenas 3 a 5 incluyen las tres capturas de prueba. El auditor activa ESCAPE desde estomago.txt; ESCAPE conserva por ahora la misma evitación de obstáculos y no localiza la salida.')
d.add_page_break()
d.add_heading('Verificación del programa',1)
p(d,'La suite local contiene 30 pruebas: 13 de Linareada, 9 de Hito 2 y 8 de robustez. Comprueba hilos y exclusión, persistencia, errores de archivo, reinicio, parser decimal estricto, respuestas ERROR, fronteras de la red y conexiones concurrentes. Se usa Neuroph 2.98, la versión declarada en Maven.')
p(d,'El test intermitente de cierre se corrigió esperando la condición cliente aceptado y registrado. No se agregó una espera fija de 300 ms. Una prueba adicional repite ese cierre 30 veces. La prueba concurrente verifica seis clientes con 40 lecturas cada uno y 24 capturas únicas. GitHub Actions está configurado para Linux con JDK 21 y Windows con JDK 17; su resultado se consulta en el pull request.')
p(d,'Una trama completa inválida o excesiva recibe ERROR|TRAMA_INVALIDA y no genera captura ni orden. El servidor conserva la conexión si puede enviar la respuesta. Se rechazan exponentes, hexadecimales, signos, espacios, NaN e infinito. EOF sin LF se descarta. Los clientes deben consumir las respuestas antes de cerrar la conexión.')
d.add_heading('Pendientes del proyecto físico',1)
p(d,'El usuario autorizó trabajar con valores ficticios. La consigna docente exige además lecturas reales y porcentaje de éxito en campo: esa evidencia sigue pendiente y no puede reemplazarse con estos resultados. También faltan Bluetooth, Safe Mode físico, identificación de las categorías finales, búsqueda de salida y comunicación de estados. El DCU en Word puede importarse a Google Docs; no se publicó un Google Doc desde esta tarea.')
d.add_heading('Archivos para comprobar los resultados',1)
p(d,'evidencias/metricas-modelo.json contiene los conteos; evidencias/predicciones.csv, cada vector, etiqueta y predicción. src/main/java/orquitas/simulacion/EvaluadorDemostracion.java reproduce la evaluación. modelos/entorno-demo.nnet es el modelo usado. El código de prueba está en src/test/java/orquitas/verificacion.')
p(d,'SHA 256 del modelo')
para=p(d,M['modelo_sha256']);para.runs[0].font.name='Consolas';para.runs[0].font.size=Pt(9)
p(d,'Referencias de alcance: HITO 2 Proyecto Orquitas La Orquita Pensante, página 1; Proyecto Orquitas Reglas Técnicas de Programación Sobre Redes, páginas 1 y 2. Los datos numéricos de este informe provienen exclusivamente de las ejecuciones locales registradas.')
d.save(R/'docs/Informe_simulacion_Hito2.docx')

# Editable DCU of the current implementation, without publishing to an external account.
d=base('Casos de uso del servidor Orquitas','Hito 2 preliminar con telemetría y datos ficticios')
text=(R/'docs/DCU_Hito2.md').read_text(encoding='utf-8')
for line in text.splitlines()[2:]:
 if not line.strip():continue
 if line.startswith('## '):d.add_heading(re.sub(r'[^\w\sáéíóúñÁÉÍÓÚÑ]',' ',line[3:]).strip(),1)
 else:
  para=d.add_paragraph(re.sub(r'[*`]+','',line))
  if re.match(r'^\d+\.',line):para.paragraph_format.left_indent=Inches(.12)
d.save(R/'docs/DCU_Hito2.docx')
print('Created two DOCX files')
